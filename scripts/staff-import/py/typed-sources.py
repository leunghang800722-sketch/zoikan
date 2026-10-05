# WP-STAFF-FIELDS：讀打字名單（Maxwell 前線名單、U Team 名單、BOCHK Payroll Editor）→ data/2026-10/staff-fields/typed-sources.json
# 只在本機 data/ 輸出；console 只印計數。用 tools/importer/venv（openpyxl／msoffcrypto／xlrd 現成）。
import io, json, re, sys, datetime
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[3]
OUT = ROOT / 'data/2026-10/staff-fields/typed-sources.json'
MSG = ROOT / 'data/2026-09/msg-file'
PE_PASSWORD = 'VelvetSweatshop'  # Excel 預設唯讀密碼（同 tools/importer 做法一致）


def iso(v):
    if v is None or v == '':
        return None
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime('%Y-%m-%d')
    return str(v).strip() or None


def s(v):
    if v is None:
        return None
    if isinstance(v, float) and v == int(v):
        v = int(v)
    return str(v).strip() or None


def hdr_index(header, *needles):
    for i, h in enumerate(header):
        t = re.sub(r'\s+', ' ', str(h or '')).lower()
        if all(n.lower() in t for n in needles):
            return i
    return None


def maxwell():
    wb = openpyxl.load_workbook(MSG / 'Maxwell All Staff List (Front Line).xlsx', read_only=True, data_only=True)
    rows = []
    for ws in wb.worksheets[:2]:
        data = list(ws.iter_rows(values_only=True))
        header = data[1]
        ix = {
            'staffNo': hdr_index(header, 'staff no'), 'nameEn': hdr_index(header, 'eng name'), 'nameZh': hdr_index(header, 'chinese name'),
            'hkid': hdr_index(header, 'hkid'), 'join': hdr_index(header, 'commencement'), 'payee': hdr_index(header, 'payee'),
            'bankCode': hdr_index(header, 'bank', 'code'), 'acct': hdr_index(header, 'bank account'),
            'sppNo': hdr_index(header, 'spp license'), 'sppA': hdr_index(header, 'spp a'), 'sppB': hdr_index(header, 'spp b'),
            'qas': hdr_index(header, 'qas'), 'green': hdr_index(header, 'green'),
        }
        em = [i for i, h in enumerate(header) if 'emergency' in str(h or '').lower()]
        for r, row in enumerate(data[2:], start=3):
            g = lambda k: row[ix[k]] if ix[k] is not None and ix[k] < len(row) else None
            if not (s(g('staffNo')) or s(g('hkid')) or s(g('nameEn'))):
                continue
            rows.append({
                'ref': f'maxwell:{ws.title}:{r}', 'sheet': ws.title, 'staffNo': s(g('staffNo')), 'nameEn': s(g('nameEn')), 'nameZh': s(g('nameZh')),
                'hkid': s(g('hkid')), 'join': iso(g('join')), 'payee': s(g('payee')), 'bankCode': s(g('bankCode')), 'acct': s(g('acct')),
                'sppNo': s(g('sppNo')), 'sppA': iso(g('sppA')), 'sppB': iso(g('sppB')), 'qas': iso(g('qas')), 'green': iso(g('green')),
                'emergency': [s(row[i]) if i < len(row) else None for i in em[:3]],
            })
    return rows


def uteam():
    wb = openpyxl.load_workbook(MSG / 'U Team - People template 20260824.xlsx', read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = []
    for r, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        g = lambda i: row[i] if i < len(row) else None
        no = s(g(13)) or s(g(0))
        if not (no or s(g(15))):
            continue
        rows.append({
            'ref': f'uteam:{r}', 'staffNo': no, 'join': iso(g(1)), 'status': s(g(4)), 'hkid': s(g(14)), 'nameZh': s(g(15)),
            'nameEn': ' '.join(x for x in [s(g(16)), s(g(17))] if x) or None,
            'emergency': [s(g(20)), s(g(21)), s(g(22))], 'bankCode': s(g(23)), 'acct': s(g(25)),
        })
    return rows


def parttime():
    # 「在勤兼職資料.xlsx 的副本」係普通 xlsx（openpyxl 淨係因副檔名「的副本」拒讀）；內容同 在勤兼職資料-copy.xlsx 逐位元相同，讀後者
    wb = openpyxl.load_workbook(MSG / '在勤兼職資料-copy.xlsx', read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = []
    for r, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        g = lambda i: row[i] if i < len(row) else None
        if not s(g(0)):
            continue
        rows.append({'ref': f'parttime:{r}', 'staffNo': s(g(0)), 'hkid': s(g(3)), 'firstStamp': iso(g(7)), 'lastStamp': iso(g(8))})
    return rows


def payroll_editor():
    import msoffcrypto, xlrd
    out = []
    for month in ('2026-07', '2026-08'):
        p = ROOT / f'data/{month}/reference/{month}-長工autopay-BOCHK-PayrollEditor.xls'
        with open(p, 'rb') as fh:
            of = msoffcrypto.OfficeFile(fh)
            of.load_key(password=PE_PASSWORD)
            buf = io.BytesIO()
            of.decrypt(buf)
        sh = xlrd.open_workbook(file_contents=buf.getvalue()).sheet_by_name('Payroll Editor')
        r = 16
        while r < sh.nrows and sh.cell_value(r, 0) not in ('', None):
            name = s(sh.cell_value(r, 4))
            acct = s(sh.cell_value(r, 3))
            if name and acct:
                out.append({'month': month, 'row': r + 1, 'bankCode': s(sh.cell_value(r, 1)), 'acctType': s(sh.cell_value(r, 2)), 'acct': acct, 'name': name})
            r += 1
    return out


if __name__ == '__main__':
    d = {'maxwell': maxwell(), 'uteam': uteam(), 'payrollEditor': payroll_editor(), 'parttime': parttime()}
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(d, ensure_ascii=False))
    print({k: len(v) for k, v in d.items()})
