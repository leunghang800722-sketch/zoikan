"""WP3（⑤ rates + ⑥ report）單元測試。

覆蓋範圍（P0-執行計劃.md §7 WP3 驗收準則 + 品質段要求）：
  - 發票 PDF 解析（parse_invoice_pdf）
  - rate 反推：唯一解 / 兩個價分解 / OT 按鐘分解 / K-1331 線性解嘅共線性偵測
  - pricing 配對：guard 專屬 rate 唔會漏俾第二個唔相干安員（真實出現過嘅
    regression bug：G0039 autopay 錯誤襲用咗 G0127 嘅 guard_wage，PDPO：
    只記guard_id）
  - monthly report：monthly_fixed 唔可以逐更乘大（真實出現過嘅 regression
    bug：IFC 35,000 x 40 條 shift_record = 1,400,000）
  - end-to-end：`run_wp3()` 喺有真實資料嘅機器上要順利跑完，關鍵不變式成立
    （冇真實資料嘅環境自動 skip，唔會 fail CI，跟 test_end_to_end.py 同一
    原則）
"""
from decimal import Decimal

import pytest

from importer import paths, wp3_pricing, wp3_rates
from importer.wp3_rates import RateCardBuilder


# ---------------------------------------------------------------------------
# 發票 PDF 解析
# ---------------------------------------------------------------------------

pytestmark_pdf = pytest.mark.skipif(
    not paths.INVOICE_PDF.exists(),
    reason="data/reference-global/Invoice-Record-*.pdf 唔存在（本機專用資料，冇入 git）",
)


@pytestmark_pdf
def test_parse_invoice_pdf_july_count_and_quoted_invoice_no():
    from importer import wp3_io
    lines = wp3_io.parse_invoice_pdf()
    jul = [l for l in lines if l["period"] == "2026-07"]
    assert len(jul) == 48
    # MW-2026-07009R2 前面有個 PDF 排版殘留嘅單引號，regex 要照樣攞到完整
    # invoice_no（唔可以斷咗喺 "-" 度）。
    r2 = [l for l in jul if "07009" in l["invoice_no"]]
    assert len(r2) == 1
    assert r2[0]["invoice_no"] == "MW-2026-07009R2"
    assert r2[0]["label"] == "HKU C更"
    assert r2[0]["amount"] == "139978.80"


# ---------------------------------------------------------------------------
# 兩個價分解 / OT 按鐘分解
# ---------------------------------------------------------------------------

def test_brute_force_two_rate_unique_solution():
    # 31 更 @900 + 33 更 @850 = 27,900 + 28,050 = 55,950
    amount = Decimal("55950")
    sols = wp3_rates._brute_force_two_rate([31, 33], amount)
    assert (Decimal("900"), Decimal("850")) in sols


def test_brute_force_two_rate_prefers_known_rates_when_ambiguous():
    # 刻意揀一個喺網格內有多過一組解嘅金額，但只有一組嘅兩個價都喺
    # known_rates 出現過 -> 應該揀嗰組，唔係亂揀。
    # 20更@700 + 10更@600 = 14000+6000=20000；20更@650+10更@700=13000+7000=20000
    # （兩組都啱，700/600 vs 650/700，用known_rates篩走唔喺歷史出現過嘅果組）
    amount = Decimal("20000")
    known = {Decimal("700"), Decimal("600")}
    sols = wp3_rates._brute_force_two_rate([20, 10], amount, known_rates=known)
    assert all(r0 in known and r1 in known for r0, r1 in sols)


def test_ot_hourly_decomposition_hku_c_shift_style():
    # 199更 @700 + 8小時 OT @(700/8.25) 四捨五入到分 = 139,978.79，同真實
    # 發票 139,978.80 只差 1 分（PDF 本身四捨五入），容許 0.02 誤差。
    amount = Decimal("139978.80")
    sol, candidates = wp3_rates._try_ot_hourly_decomposition(
        amount, total_shifts=198, known_rates={Decimal("700")})
    assert sol is not None
    count, rate, ot_hours = sol
    assert count == 199
    assert rate == Decimal("700")
    assert ot_hours == Decimal("8")


def test_ot_hourly_decomposition_returns_none_when_truly_ambiguous():
    # 冇 known_rates 提示、又有多過一組合理解 -> 唔可以屈一個出嚟。
    sol, candidates = wp3_rates._try_ot_hourly_decomposition(
        Decimal("139978.80"), total_shifts=198, known_rates=None)
    assert sol is None
    assert len(candidates) > 1


# ---------------------------------------------------------------------------
# pricing 配對：guard 專屬 rate 唔會漏俾第二個唔相干安員
# ---------------------------------------------------------------------------

def _fake_broker_index():
    from importer.normalize import BrokerIndex
    return BrokerIndex([
        {"broker_id": "B05", "code": "T", "canonical_name": "T",
         "aliases": [], "kind": "broker", "settlement_type": "cash"},
    ])


def test_pricing_does_not_leak_rate_across_guards_with_same_site_broker():
    """Regression test：創富車廠(夜) T 頁入面，G0039（autopay，
    guard_wage=None, broker_fee=30）同 G0127（現金，guard_wage=680,
    broker_fee=50）係同一個 (site,broker) 但完全唔同人、唔同價（PDPO：呢度
    只記 guard_id，具體人可查 out/wp2b/guards.csv）。之前個
    bug：pricing.py 揾 G0039 自己嗰張卡搵唔到 guard_wage（佢本身冇），就
    再向粗嘅 (site,broker) 層搵，喺嗰層撞到 G0127 張卡，錯誤攞咗680嚟用。
    修正之後：G0039 查 guard_wage 應該係 None（唔會屈用第二個人嘅價），
    查 broker_fee 應該啱（佢自己嘅30）。下面 fixture 用 G_LI/G_WONG 呢啲
    假 guard_id，同真實 guard_id 冇關係。"""
    rcb = RateCardBuilder()
    rcb.add(site_id="S019", broker_id="B05", guard_id="G_LI",
             guard_wage=None, broker_fee=Decimal("30"), payout_rate=Decimal("30"),
             billing_model="per_shift", source="test", evidence="", confidence="high",
             status="solved")
    rcb.add(site_id="S019", broker_id="B05", guard_id="G_WONG",
             guard_wage=Decimal("680"), broker_fee=Decimal("50"), payout_rate=Decimal("730"),
             billing_model="per_shift", source="test", evidence="", confidence="high",
             status="solved")
    lookup = wp3_pricing.build_rate_lookup(rcb.rows)

    li_wage = wp3_pricing._pick_card(lookup, "S019", "ST_ANY", "B05", "G_LI", "guard_wage")
    li_fee = wp3_pricing._pick_card(lookup, "S019", "ST_ANY", "B05", "G_LI", "broker_fee")
    assert li_wage is None, "G0039自己冇guard_wage，唔應該攞到G0127嘅680"
    assert li_fee is not None and li_fee["broker_fee"] == "30.00"

    wong_wage = wp3_pricing._pick_card(lookup, "S019", "ST_ANY", "B05", "G_WONG", "guard_wage")
    assert wong_wage is not None and wong_wage["guard_wage"] == "680.00"


def test_pricing_falls_back_to_site_broker_when_guard_has_no_own_card():
    """一個冇專屬card嘅第三個安員（例如guard_id resolve失敗），一樣可以
    fall back 用 (site,broker) 層嘅價——但如果嗰層有多過一個唔同價就要
    ambiguous（None），唔可以隨機揀一個。"""
    rcb = RateCardBuilder()
    rcb.add(site_id="S020", broker_id="B05", guard_id="",
             guard_wage=Decimal("820"), broker_fee=None, payout_rate=Decimal("820"),
             billing_model="per_shift", source="test", evidence="", confidence="high",
             status="solved")
    lookup = wp3_pricing.build_rate_lookup(rcb.rows)
    picked = wp3_pricing._pick_card(lookup, "S020", "ST_X", "B05", "", "guard_wage")
    assert picked is not None and picked["guard_wage"] == "820.00"


# ---------------------------------------------------------------------------
# WP5g：master_roster_payout（A4）出街價最後 fallback——單元測試（唔靠真實
# 資料）。regression test：呢個 fallback 一定要淨係喺主 cascade（糧明細/
# master_roster_rate/master_roster_only_rate）完全搵唔到 candidate 先至
# apply，唔可以打橫補一個已經有「真卡但特登留空」嘅欄位（例如長工autopay
# 喺 T 頁淨收介紹費，guard_wage 特登空、broker_fee 有值嗰種卡）。
# ---------------------------------------------------------------------------

def _rec_for_payout(shift_id, guard_id, site_id, shift_type_id, settlement_unit):
    return {
        "shift_id": shift_id, "guard_id": guard_id, "site_id": site_id,
        "shift_type_id": shift_type_id, "settlement_unit": settlement_unit,
        "is_external": "False", "out_of_scope": "False",
        "record_type": "shift", "work_date": "2026-07-03",
        "referral_broker": "",
    }


def test_master_roster_payout_fallback_fills_guard_with_no_primary_card():
    """主cascade完全冇呢個guard/site/broker嘅卡 -> 應該用返
    master_roster_payout 嘅總表長工卡（guard_wage=840.38, broker_fee=0，
    總表外判價全包）。"""
    from importer.normalize import BrokerIndex
    broker_index = BrokerIndex([
        {"broker_id": "B06", "code": "Y", "canonical_name": "公司兵",
         "aliases": [], "kind": "internal_cash", "settlement_type": "cash"},
    ])
    rcb = RateCardBuilder()
    rcb.add(site_id="S017", broker_id="B06", guard_id="G_TEST_A",
            guard_wage=Decimal("840.38"), broker_fee=Decimal("0"), payout_rate=Decimal("840.38"),
            billing_model="per_shift", source="master_roster_payout", evidence="", confidence="high",
            status="solved")
    priced, unpriced = wp3_pricing.price_shift_records(
        [_rec_for_payout("SR_A", "G_TEST_A", "S017", "ST0026", "公司兵")],
        rcb.rows, broker_index, {}, shift_type_billing_group={})
    assert not unpriced
    assert priced[0]["guard_wage_snapshot"] == "840.38"
    assert priced[0]["broker_fee_snapshot"] == "0.00"


def test_build_master_roster_payout_rate_cards_two_guards_share_site_broker():
    """Regression test（WP5g 開發期間實測過嘅真實 bug）：兩個長工共用同一個
    (site,broker)——例如啟航1331 早更保安員 桶Y公司兵入面 G9001/G9002
    兩位，各自喺 master_posts 有專屬釘死嘅崗位、各自唔同嘅
    payout_rate（840.38 vs 718.02）。如果 `_already_covered` check 用緊
    「不斷變化」嘅 rcb.rows（每加一張新卡就即刻影響落一個 check），第一個
    guard（G9001）張卡一加咗，佢個 broker_fee=0（總表卡設計）會令
    (site,broker) 呢層闊key（site_br）由「多個candidate冇統一值」變成
    「得返一個非空broker_fee candidate」，令第二個guard（G9002）嘅
    _already_covered check 誤判「已經有值」，令佢自己張卡冇生成。修正之後
    （`_already_covered` 用一個入 function 之前嘅 rcb.rows 快照，唔會隨住
    function 自己逐張加卡而變動）：兩個guard都應該有自己嘅卡。"""
    from importer.normalize import BrokerIndex
    broker_index = BrokerIndex([
        {"broker_id": "B06", "code": "Y", "canonical_name": "公司兵",
         "aliases": [], "kind": "internal_cash", "settlement_type": "cash"},
    ])
    shift_types = [{"site_id": "S_TEST", "shift_type_id": "ST_TEST",
                     "time_slot": "0730-1930", "billing_group": "per_shift"}]
    posts = [
        {"post_id": "MPT1", "site_id": "S_TEST", "shift": "Day", "time": "0730-1930",
         "payout_rate": "840.38", "guard_name_hint": "[NAME:G9001]", "customer_label": ""},
        {"post_id": "MPT2", "site_id": "S_TEST", "shift": "Day", "time": "0730-1930",
         "payout_rate": "718.02", "guard_name_hint": "[NAME:G9002]", "customer_label": ""},
    ]
    assignments = (
        [{"post_id": "MPT1", "broker_code": "Y"} for _ in range(26)]
        + [{"post_id": "MPT2", "broker_code": "Y"} for _ in range(26)]
    )
    rcb = RateCardBuilder()
    meta = wp3_rates.build_master_roster_payout_rate_cards(
        posts, assignments, shift_types, broker_index, rcb)
    assert meta["guard_cards"] == 2, f"兩個guard都應該有自己嘅總表長工卡：{meta}"
    by_guard = {rc["guard_id"]: rc for rc in rcb.rows if rc.get("guard_id")}
    assert by_guard["G9001"]["guard_wage"] == "840.38"
    assert by_guard["G9002"]["guard_wage"] == "718.02", \
        "第二個guard唔應該因為第一個guard張卡岩岩好帶咗個非空broker_fee=0而被誤判「已有值」"


def test_master_roster_payout_fallback_does_not_override_referral_only_card():
    """Regression test（WP5g 開發期間實測過嘅真實 bug）：長工 autopay 喺
    T 頁淨收介紹費（guard_wage 特登留空、broker_fee=30 有值——同
    G0039 個 pattern 一致），呢種「已經有真卡，只不過呢個field特登冇值」
    嘅記錄，master_roster_payout fallback 唔可以打橫補一個唔啱嘅
    guard_wage（之前個bug：某長工T頁淨收780介紹費，guard_wage特登
    空，fallback錯誤補咗成月19,760）。修正之後：fallback 淨係喺 wage_rc
    *同* fee_rc 都係 None（乜都搵唔到）先至apply，呢度 fee_rc 已經搵到
    （介紹費30），wage_rc 應該維持 None，唔會被 fallback 補。"""
    from importer.normalize import BrokerIndex
    broker_index = BrokerIndex([
        {"broker_id": "B05", "code": "T", "canonical_name": "T",
         "aliases": [], "kind": "broker", "settlement_type": "cash"},
    ])
    rcb = RateCardBuilder()
    # 主cascade：呢個guard自己嘅「介紹費」卡（guard_wage留空，broker_fee=30）
    rcb.add(site_id="S019", broker_id="B05", guard_id="G_TEST_B",
            guard_wage=None, broker_fee=Decimal("30"), payout_rate=Decimal("30"),
            billing_model="per_shift", source="payroll-rows.csv", evidence="", confidence="high",
            status="solved")
    # master_roster_payout fallback：同一個(site,broker)嘅崗位卡（唔應該
    # 打橫補呢個guard嘅guard_wage，因為佢自己已經有張真卡）
    rcb.add(site_id="S019", shift_type_id="ST_ANY", broker_id="B05",
            guard_wage=Decimal("700"), broker_fee=Decimal("0"), payout_rate=Decimal("700"),
            billing_model="per_shift", source="master_roster_payout", evidence="", confidence="high",
            status="solved")
    priced, unpriced = wp3_pricing.price_shift_records(
        [_rec_for_payout("SR_B", "G_TEST_B", "S019", "ST_ANY", "T")],
        rcb.rows, broker_index, {}, shift_type_billing_group={})
    assert not unpriced
    assert priced[0]["guard_wage_snapshot"] == "", \
        "呢個guard自己張卡特登冇guard_wage（淨收介紹費），fallback唔應該補一個假wage"
    assert priced[0]["broker_fee_snapshot"] == "30.00"


# ---------------------------------------------------------------------------
# WP3 第三版 A3：special 卡（帶 shift_type_id）唔可以喺 site 層洩漏俾
# 其他冇專屬卡嘅 shift_type（真實出現過嘅 regression bug：啟航1331 Event
# RC0153 client_rate=11,700 洩漏俾 ST0027/ST0029 共 248 條非 Event 記錄）
# ---------------------------------------------------------------------------

def test_special_card_with_shift_type_does_not_leak_to_site_level_key():
    """`_rc_key_candidates` 一旦張卡帶 shift_type_id，就唔應該再喺
    ("site", site) 呢層粗 key 出現——否則冇專屬卡嘅其他 shift_type 會經
    site fallback 撞用到呢張本身淨係俾一個 shift_type 用嘅 special 卡。"""
    rcb = RateCardBuilder()
    rcb.add(site_id="S017", shift_type_id="ST0026", client_rate=Decimal("11700"),
            billing_model="special", source="test", evidence="", confidence="high",
            status="solved")
    lookup = wp3_pricing.build_rate_lookup(rcb.rows)
    assert ("site", "S017") not in lookup, \
        "special 卡帶 shift_type_id 就唔應該登記喺 (\"site\", site) 呢層"
    assert ("site_st", "S017", "ST0026") in lookup


def test_special_card_price_only_applies_to_its_own_shift_type():
    """端對端（`price_shift_records`）版本：special 卡（ST0026 Event，
    client_rate=11,700）同一個 site 入面，另一個冇任何 rate_card 嘅
    shift_type（ST0027）唔可以撞用到 11,700；ST0026 自己嘅記錄要繼續攞
    到（client_rate_id 指返呢張卡，special 制唔逐更 stamp snapshot）。"""
    from importer.normalize import BrokerIndex

    rcb = RateCardBuilder()
    rcb.add(site_id="S017", shift_type_id="ST0026", client_rate=Decimal("11700"),
            billing_model="special", source="test", evidence="", confidence="high",
            status="solved")
    broker_index = BrokerIndex([])

    def _rec(shift_id, shift_type_id):
        return {
            "shift_id": shift_id, "guard_id": "", "site_id": "S017",
            "shift_type_id": shift_type_id, "settlement_unit": "",
            "is_external": "False", "out_of_scope": "False",
            "record_type": "shift", "work_date": "2026-07-03",
            "referral_broker": "",
        }

    priced, unpriced = wp3_pricing.price_shift_records(
        [_rec("SR_EVENT", "ST0026"), _rec("SR_OTHER", "ST0027")],
        rcb.rows, broker_index, {}, shift_type_billing_group={})

    priced_by_id = {r["shift_id"]: r for r in priced}
    assert "SR_EVENT" in priced_by_id
    assert priced_by_id["SR_EVENT"]["client_rate_id"] == "RC0001"
    assert priced_by_id["SR_EVENT"]["client_rate_snapshot"] == "", \
        "special 制唔逐更 stamp client_rate_snapshot，淨係留 client_rate_id"

    unpriced_ids = {r["shift_id"] for r in unpriced}
    assert "SR_OTHER" in unpriced_ids, \
        "ST0027 冇自己嘅 rate_card，唔應該撞用 ST0026 Event 嗰張 special 卡"
    other_by_id = {r["shift_id"]: r for r in unpriced}
    assert "11700" not in other_by_id["SR_OTHER"]["reason"]


# ---------------------------------------------------------------------------
# monthly report：monthly_fixed 唔可以逐更乘大
# ---------------------------------------------------------------------------

@pytest.mark.skipif(
    not paths.SETTLEMENT_PDF_2026_07.exists(),
    reason="data/2026-07/reference/2026-07-月結總表-FinalR2.pdf 唔存在（本機專用資料，冇入 git）",
)
def test_monthly_report_does_not_multiply_monthly_fixed_by_shift_count(tmp_path):
    from importer import wp3_reports

    rate_cards = [{
        "rate_id": "RC1", "site_id": "S002", "shift_type_id": "", "broker_id": "",
        "guard_id": "", "client_rate": "35000.00", "guard_wage": "", "broker_fee": "",
        "payout_rate": "", "billing_model": "monthly_fixed", "effective_from": "2026-07-01",
        "effective_to": "2026-07-31", "source": "test", "evidence": "", "confidence": "high",
        "status": "solved", "note": "",
    }]
    priced_rows = [
        {"site_id": "S002", "record_type": "shift", "shift_type_id": "ST_TEST",
         "is_external": "False", "out_of_scope": "False", "pricing_note": "",
         "client_rate_snapshot": "", "guard_wage_snapshot": "", "broker_fee_snapshot": ""}
        for _ in range(40)
    ]
    sites_by_id = {"S002": {"site_id": "S002", "canonical_name": "IFC",
                             "billing_model": "monthly_fixed", "client_name": "",
                             "aliases": [], "client_group": ""}}

    monthly_diff, settle_rows = wp3_reports.build_monthly_report(
        tmp_path, priced_rows, sites_by_id, invoice_aliases={}, rate_cards=rate_cards)

    # 由 workbook 讀返 IFC 嗰行嘅營業額，確認唔係 35000*40。新表頭（WP3v2
    # 條件2）：工作地點/site_id/更數/收客價/出街價/營業額/成本/盈利/%，
    # 營業額而家喺index 5（0-indexed），唔再係3。
    import openpyxl
    wb = openpyxl.load_workbook(tmp_path / "月結總表-regenerated.xlsx")
    ws = wb["月結對照(R2口徑)"]
    header = next(r for r in ws.iter_rows(values_only=True))
    revenue_col = header.index("營業額")
    ifc_rows = [r for r in ws.iter_rows(values_only=True) if r and r[1] == "S002"]
    assert len(ifc_rows) == 1
    revenue = Decimal(str(ifc_rows[0][revenue_col]))
    assert revenue == Decimal("35000"), f"monthly_fixed revenue 唔應該乘更數，攞到{revenue}"


@pytest.mark.skipif(
    not paths.SETTLEMENT_PDF_2026_07.exists(),
    reason="data/2026-07/reference/2026-07-月結總表-FinalR2.pdf 唔存在（本機專用資料，冇入 git）",
)
def test_monthly_report_shift_count_excludes_external_and_monthly_ic(tmp_path):
    """WP3v2 條件2第1點：更數要用 shift_records（record_type=shift、
    is_external=False、out_of_scope=False、billing_group=per_shift）數，
    external（客戶自己嘅人）同 monthly_ic（月費制嘅早更IC）唔可以計落
    per_shift 更數。"""
    from importer import wp3_reports

    rate_cards = [{
        "rate_id": "RC1", "site_id": "S011", "shift_type_id": "", "broker_id": "",
        "guard_id": "", "client_rate": "930.00", "guard_wage": "", "broker_fee": "",
        "payout_rate": "", "billing_model": "per_shift", "effective_from": "2026-07-01",
        "effective_to": "2026-07-31", "source": "test", "evidence": "", "confidence": "high",
        "status": "solved", "note": "",
    }]
    priced_rows = (
        [{"site_id": "S011", "record_type": "shift", "shift_type_id": "ST_NORMAL",
          "is_external": "False", "out_of_scope": "False", "pricing_note": "",
          "client_rate_snapshot": "930.00", "guard_wage_snapshot": "", "broker_fee_snapshot": ""}
         for _ in range(10)]
        + [{"site_id": "S011", "record_type": "shift", "shift_type_id": "ST_EXTERNAL",
            "is_external": "True", "out_of_scope": "False", "pricing_note": "external",
            "client_rate_snapshot": "", "guard_wage_snapshot": "", "broker_fee_snapshot": ""}
           for _ in range(3)]
        + [{"site_id": "S011", "record_type": "shift", "shift_type_id": "ST_IC",
            "is_external": "False", "out_of_scope": "False", "pricing_note": "monthly_ic",
            "client_rate_snapshot": "", "guard_wage_snapshot": "", "broker_fee_snapshot": ""}
           for _ in range(5)]
    )
    shift_type_billing_group = {"ST_NORMAL": "per_shift", "ST_EXTERNAL": "per_shift", "ST_IC": "monthly_ic"}
    sites_by_id = {"S011": {"site_id": "S011", "canonical_name": "數碼港 CP5",
                             "billing_model": "per_shift", "client_name": "", "aliases": [], "client_group": ""}}

    monthly_diff, settle_rows = wp3_reports.build_monthly_report(
        tmp_path, priced_rows, sites_by_id, invoice_aliases={}, rate_cards=rate_cards,
        shift_type_billing_group=shift_type_billing_group)

    import openpyxl
    wb = openpyxl.load_workbook(tmp_path / "月結總表-regenerated.xlsx")
    ws = wb["月結對照(R2口徑)"]
    header = next(r for r in ws.iter_rows(values_only=True))
    shifts_col = header.index("更數")
    cp5_rows = [r for r in ws.iter_rows(values_only=True) if r and r[1] == "S011"]
    assert len(cp5_rows) == 1
    shifts_display = str(cp5_rows[0][shifts_col])
    assert shifts_display.startswith("10"), (
        f"更數應該淨係計10個per_shift正常記錄（external 3個、monthly_ic 5個都唔計落主更數），攞到{shifts_display}")
    assert "monthly_ic 5" in shifts_display, "monthly_ic 更數要另列，唔可以直接消失"


# ---------------------------------------------------------------------------
# end-to-end（要真實 out/wp2b/* + data/，冇就 skip）
# ---------------------------------------------------------------------------

pytestmark_e2e = pytest.mark.skipif(
    not paths.WP2B_SHIFT_RECORDS_CSV.exists() or not paths.INVOICE_PDF.exists(),
    reason="out/wp2b/*.csv 或發票 PDF 唔存在（要先行 run-wp2b，本機專用資料）",
)


@pytestmark_e2e
def test_run_wp3_end_to_end_invariants():
    from importer.wp3_run import run_wp3
    summary = run_wp3()
    assert summary["rate_cards"] > 0
    assert summary["solved"] > 0
    # 每個 rate_card 要麼 solved 要麼 pending，唔會兩個都唔係
    assert summary["solved"] + summary["pending"] == summary["rate_cards"]
    # 抽樣覆核全部應該 match（見 sample-audit.csv）
    import csv
    with open(paths.WP3_SAMPLE_AUDIT_CSV, encoding="utf-8") as fh:
        rows = list(csv.DictReader(fh))
    assert rows, "sample-audit.csv 唔應該係空"
    bad = [r for r in rows if r["match"] != "TRUE"]
    # 已知、已喺 wp3_rates.KNOWN_SHIFT_COUNT_DIFFS / billing_units.csv 記低
    # 出處嘅 site（例如HKU C更 S007：OT-hourly解用嘅更數199同shift_records
    # 實際更數198差1；控制室 S031：R2話66更但shift_records淨係得44更，屬於
    # WP1/WP2b層面已知嘅資料流缺口）容許喺呢個粗略嘅billing_unit_sample
    # 隨機抽樣入面顯示REVIEW——真正嘅收貨標準係wp3_rates.apply_client_rate_
    # gate 嘅 gate-results.csv 全部 PASS（見下面 test_client_rate_gate_all_pass），
    # 呢度淨係抽樣層面唔可以再有第二類、未解釋嘅唔match。
    known_gap_site_ids = set(wp3_rates.KNOWN_SHIFT_COUNT_DIFFS)
    # P6 已知嘅細殘差（billing_adjustments.csv，例如西九龍城巴車廠+30、
    # 控制室+1,350）都係已經逐條記低出處、等老闆一次過確認嘅已知項，隨機
    # 抽樣揀中都唔算「未解釋」。
    with open(paths.WP3_BILLING_ADJUSTMENTS_CSV, encoding="utf-8") as fh:
        known_gap_site_ids |= {r["site_id"] for r in csv.DictReader(fh)}
    unexpected_bad = [r for r in bad
                       if not any(f"site={sid}" in r["key"] or f",{sid}," in r["key"]
                                  for sid in known_gap_site_ids)]
    assert not unexpected_bad, f"抽樣覆核有未解釋嘅唔match：{unexpected_bad}"
    # roundtrip diff 淨係應該有已知類別（external_row_excluded / identity_*），
    # 唔應該有 needs_review（呢個表示WP1->WP2b->WP3之間漏咗格仔）
    with open(paths.WP3_ROUNDTRIP_DIFF_CSV, encoding="utf-8") as fh:
        rt = list(csv.DictReader(fh))
    bad_categories = [r for r in rt if r["category"] == "needs_review"]
    assert not bad_categories, f"round-trip 出現未解釋嘅needs_review：{bad_categories}"


# ---------------------------------------------------------------------------
# WP3 第三版 B12：Gate 負向測試（單元級，唔靠真實資料）——注入一個錯價，
# 斷言 apply_client_rate_gate 會判 FAIL 並將對應 rate_card 降返 pending。
# 之前只有 test_client_rate_gate_all_pass（正向：真實資料全部PASS），冇
# 一個測試證明個Gate真係「有牙」——會唔會其實成個check冇運作、淨係啱啱好
# 真實資料全部啱先睇落好似work緊。
# ---------------------------------------------------------------------------

def test_client_rate_gate_fails_and_downgrades_on_wrong_price():
    rcb = RateCardBuilder()
    rcb.add(site_id="S_GATE_TEST", broker_id="B_TEST", client_rate=Decimal("100"),
            billing_model="per_shift", source="test", evidence="", confidence="high",
            status="solved")
    rate_cards = rcb.rows

    # 逐更 client_rate_snapshot 全部100，10條記錄，Σ=1000
    priced_rows = [
        {"site_id": "S_GATE_TEST", "record_type": "shift",
         "client_rate_snapshot": "100.00"}
        for _ in range(10)
    ]

    # 錯價注入：發票話呢個site個月收咗5000（同Σ1000差成4000，遠超±0.05
    # 容忍度，亦冇任何已知shift-count差／billing_adjustments可以cover）
    client_solve_log = [{
        "invoice_no": "INV-TEST-001", "label": "假盤GateTest", "site_id": "S_GATE_TEST",
        "amount": "5000.00", "shifts": "10", "client_rate": "100",
        "status": "solved(test_injected_wrong_price)", "note": "",
    }]

    gate_rows, downgraded_sites = wp3_rates.apply_client_rate_gate(
        rate_cards, priced_rows, client_solve_log, billing_adjustments=[])

    assert "S_GATE_TEST" in downgraded_sites, "錯價注入應該令個site落入downgraded_sites"
    fail_rows = [g for g in gate_rows if g["site_id"] == "S_GATE_TEST"]
    assert fail_rows and all(g["gate"] == "FAIL" for g in fail_rows), \
        f"Σclient_rate_snapshot(1000) 同發票(5000)差4000，遠超容忍度，Gate應該FAIL：{fail_rows}"

    downgraded_cards = [rc for rc in rate_cards if rc["site_id"] == "S_GATE_TEST"]
    assert downgraded_cards and all(rc["status"] == "pending" for rc in downgraded_cards), \
        "Gate FAIL 之後，對應嘅solved per_shift rate_card應該被mutate做pending"
    assert all("Gate唔過" in rc["note"] for rc in downgraded_cards), \
        "降級嘅rate_card嘅note要留低Gate唔過嘅痕跡，唔可以靜靜降級"


@pytestmark_e2e
def test_client_rate_gate_all_pass():
    """P0-執行計劃 WP3v2 條件1第3點 Gate invariant：任何 status 帶「solved」
    嘅 billing-unit-solve.csv 記錄，Σclient_rate_snapshot（+billing_
    adjustments，+已知shift-count差容忍度）都要同發票金額打平（見
    `wp3_rates.apply_client_rate_gate`），否則 run_wp3() 自己會將佢降返
    pending——所以呢度嘅不變式係「跑完之後 gate-results.csv 冇任何一條
    FAIL」，同時「rate_cards.csv 冇任何一條 status=solved 嘅 per_shift 卡
    對應住一個 FAIL 咗嘅 billing unit」（雙重check：Gate機制本身冇被
    跳過）。"""
    import csv
    from importer.wp3_run import run_wp3
    run_wp3()
    with open(paths.WP3_GATE_RESULTS_CSV, encoding="utf-8") as fh:
        gate_rows = list(csv.DictReader(fh))
    assert gate_rows, "gate-results.csv 唔應該係空"
    fails = [g for g in gate_rows if g["gate"] == "FAIL"]
    assert not fails, f"Gate 有未解決嘅 FAIL（Σ vs 發票唔啱，理應已經自動降返pending）：{fails}"

    with open(paths.WP3_BILLING_UNIT_SOLVE_CSV, encoding="utf-8") as fh:
        solve_rows = list(csv.DictReader(fh))
    gate_fail_sites = {g["site_id"] for g in gate_rows if g["gate"] == "FAIL"}
    leaked = [r for r in solve_rows
              if r["site_id"] in gate_fail_sites and r["status"].startswith("solved")]
    assert not leaked, f"呢啲site Gate FAIL 但 billing-unit-solve.csv 仲話solved：{leaked}"


@pytestmark_e2e
def test_s017_kai1331_main_invoice_gate_real_pass():
    """WP5c必修6（P0-執行計劃.md §8「WP5總表對數後新問題」item 6）：啟航1331
    主發票（MW-2026-07005，1,203,900）以前生成一張永遠pending嘅RC0155（發票
    ÷全部shift_type混埋嘅1,126更、唔整除；兩價分解都解唔到）。而家用真口徑
    （Σ6個已solved非Event更別client_rate_snapshot + 4個Event Guard(1000-1900)
    cell以1,200/更計入主發票 + 已知1更count缺口）驗證，rate_cards.csv 唔應該
    再有RC0155呢類pending卡，gate-results.csv 要有S017呢行、gate=PASS
    （唔係PASS(known_shift_gap)，係真正diff=0.00）。"""
    import csv
    from importer.wp3_run import run_wp3
    run_wp3()

    with open(paths.WP3_RATE_CARDS_CSV, encoding="utf-8") as fh:
        rate_cards = list(csv.DictReader(fh))
    pending = [rc for rc in rate_cards if rc["status"] == "pending"]
    assert not pending, f"WP5c必修6之後rate_cards.csv唔應該再有pending卡：{pending}"

    with open(paths.WP3_GATE_RESULTS_CSV, encoding="utf-8") as fh:
        gate_rows = list(csv.DictReader(fh))
    s017_rows = [g for g in gate_rows if g["site_id"] == "S017"]
    assert len(s017_rows) == 1, f"S017 應該淨係得一行gate結果：{s017_rows}"
    g = s017_rows[0]
    assert g["invoice_no"] == "MW-2026-07005"
    assert g["invoice_amount"] == "1203900.00"
    assert g["snapshot_sum"] == "1197900.00", (
        "Σ6個非Event更別已solved嘅client_rate_snapshot 預期 1,197,900"
        "（master_roster_rate卡，P0-執行計劃§8 P1已答）")
    assert g["billing_adjustments"] == "6000.00", (
        "4個Event Guard cell×1,200=4,800 + 已知1更count缺口殘差1,200 = 6,000")
    assert g["diff"] == "0.00"
    assert g["gate"] == "PASS", "真PASS（唔係PASS(known_shift_gap)）"

    with open(paths.WP3_BILLING_UNIT_SOLVE_CSV, encoding="utf-8") as fh:
        solve_rows = list(csv.DictReader(fh))
    main = next(r for r in solve_rows if r["invoice_no"] == "MW-2026-07005")
    assert main["status"] == "solved(master_roster_verified)"
    # 內部一致性：Σ非Event snapshot(1,197,900) + 已知1更缺口殘差(1,200) =
    # 1,199,100 = R2月結row21「啟德1331」印嘅營業額（獨立來源互證）。
    assert "1199100" in main["note"].replace(",", "")


@pytestmark_e2e
def test_master_roster_only_sites_real_gate_pass():
    """WP5c必修12（P0-執行計劃.md §8「WP5總表對數後新問題」item 12）：
    「月結有、更表冇」site（國泰城/ICBC/哈羅/東濤苑/控制室——M+/CP5 IC/
    威信Patrol係monthly_fixed billing_model，gate本身唔靠shift_records，
    唔喺呢個test覆蓋範圍）由總表補入庫真.逐更rate之後，gate要由
    PASS(no_shift_records)/PASS(known_shift_gap)變真PASS（或者quantize
    尾數之內）。"""
    import csv
    from importer.wp3_run import run_wp3
    run_wp3()

    with open(paths.WP2B_SHIFT_RECORDS_CSV, encoding="utf-8") as fh:
        shift_records = list(csv.DictReader(fh))
    mro_counts = {}
    for r in shift_records:
        if r["source"] == "master_roster_only" and r["record_type"] == "shift":
            mro_counts[r["site_id"]] = mro_counts.get(r["site_id"], 0) + 1
    assert mro_counts.get("S038") == 84, f"國泰城應該補178入面嘅84更：{mro_counts.get('S038')}"
    assert mro_counts.get("S034") == 25, f"M+應該補25更：{mro_counts.get('S034')}"
    assert mro_counts.get("S012") == 24, f"數碼港CP5 IC應該補24更：{mro_counts.get('S012')}"
    assert mro_counts.get("S031") == 22, f"控制室應該補22更（第3位安員）：{mro_counts.get('S031')}"
    assert mro_counts.get("S032") == 20, f"威信Patrol應該補20更：{mro_counts.get('S032')}"
    assert mro_counts.get("S041") == 2, f"ICBC Tower應該補2更：{mro_counts.get('S041')}"
    assert mro_counts.get("S042") == 1, f"哈羅國際學校應該補1更：{mro_counts.get('S042')}"
    # 東濤苑總表7/31嘅1更同G糧單頁已有記錄係同一日、同一site——去重
    # 之後應該淨係得返payroll_only果條，master_roster_only唔應該再開多
    #一條假更。
    assert "S035" not in mro_counts, (
        f"東濤苑總表嗰1格同G糧單頁重複，去重之後唔應該有master_roster_only記錄：{mro_counts}")

    with open(paths.WP3_GATE_RESULTS_CSV, encoding="utf-8") as fh:
        gate_rows = {r["site_id"]: r for r in csv.DictReader(fh)}
    for site_id, label in (("S038", "國泰城"), ("S041", "ICBC Tower"), ("S042", "哈羅國際學校")):
        g = gate_rows.get(site_id)
        assert g is not None, f"{label}應該有gate-results.csv行"
        assert g["gate"] == "PASS", f"{label}應該真PASS（diff=0.00），實際：{g}"
        assert g["diff"] == "0.00"
    control = gate_rows.get("S031")
    assert control is not None
    assert control["gate"] in ("PASS", "PASS(known_shift_gap)")
    assert abs(Decimal(control["diff"])) < Decimal("1.00"), (
        f"控制室真實snapshot之後應該只剩quantize尾數（<$1），實際：{control}")

    with open(paths.WP3_BILLING_ADJUSTMENTS_CSV, encoding="utf-8") as fh:
        adjustments = list(csv.DictReader(fh))
    assert not any(a["site_id"] == "S031" for a in adjustments), (
        "控制室嘅billing_adjustments（舊「+1,350」估計）而家由真實snapshot"
        "自己解釋咗，唔應該再有")


@pytestmark_e2e
def test_run_wp3_three_times_byte_identical_csv():
    from importer.wp3_run import run_wp3
    import hashlib

    def _hash():
        h = hashlib.md5()
        for p in sorted(paths.OUT_WP3_DIR.rglob("*.csv")):
            h.update(p.read_bytes())
        return h.hexdigest()

    run_wp3()
    h1 = _hash()
    run_wp3()
    h2 = _hash()
    run_wp3()
    h3 = _hash()
    assert h1 == h2 == h3


# ---------------------------------------------------------------------------
# WP3c：三份報表版面對齊人手版（roster round-trip / payroll / monthly-diff）
# ---------------------------------------------------------------------------

# 三種欄位排法各揀一個代表檔（P0-執行計劃§3.1）：標準17檔／數碼港3檔（多
# 英文名欄）／城巴1檔（多Property Name欄）。
_STANDARD_FILE = "Roster - 3172 殷樺花園 - Jul 2026 (updated 27.07.2026).xlsx"
_CP_FILE = "Roster - 9615 數碼港 CP5 - Jul 2026 (updated 27.07.2026).xlsx"
_CITYBUS_FILE = "Roster - 城巴車廠 - Jul 2026 (updated 27.07.2026).xlsx"


@pytestmark_e2e
def test_roster_roundtrip_needs_review_zero_per_layout():
    """三種排法（標準／數碼港／城巴）各揀一個代表檔，round-trip diff 逐檔
    needs_review 要係0——呢個係test_run_wp3_end_to_end_invariants嗰個全域
    check嘅逐檔細分版，專門對應P0-執行計劃WP3c品質要求「更表round-trip對
    一個標準檔、一個數碼港檔、城巴檔needs_review=0」。"""
    import csv
    from importer.wp3_run import run_wp3
    run_wp3()
    with open(paths.WP3_ROUNDTRIP_DIFF_CSV, encoding="utf-8") as fh:
        rt = list(csv.DictReader(fh))
    for label, file_name in (("標準", _STANDARD_FILE), ("數碼港", _CP_FILE), ("城巴", _CITYBUS_FILE)):
        bad = [r for r in rt if r["file"] == file_name and r["category"] == "needs_review"]
        assert not bad, f"{label}檔（{file_name}）round-trip 有未解釋嘅needs_review：{bad}"


@pytestmark_e2e
def test_payroll_page_totals_match_policy_total_shifts():
    """P0-執行計劃WP3c §B：糧明細逐頁總計（更數）要同
    data/2026-07/transcribed/payroll-pages.csv 嘅policy_total_shifts啱（呢
    個口徑WP2a已經驗收過）。實測T/W/公司兵/D四頁有細殘差
    （-1.25／-3／+4／+2更；WP5e應修8：公司兵由舊嘅-5修正做+4，跟返
    payroll-diff.csv__PAGE_TOTAL__嗰行即場算出嚟嘅數），全部細過每頁
    policy總更數嘅3%，而且已經喺
    payroll-diff.csv用category=shift_count_diff逐guard解釋咗（唔係
    tool_bug，見test_payroll_diff_has_no_tool_bug）——呢個殘差嚟自WP1/
    WP2b層面更表↔糧明細轉錄嘅個別guard更數唔啱（唔係WP3c report層可以
    改嘅嘢，WP3c嘅職責係診斷清楚、分類岩，唔係喺呢層砌數砌到啱）。所以
    呢個test斷言嘅係「殘差細（有解釋、冇爆數）」，唔係「殘差=0」：page_
    total mismatch 一定要有對應嘅shift_count_diff category（唔可以係
    tool_bug／未分類），並且相對policy總更數嘅比例要細（<5%），超出呢個
    範圍就當做需要WP1/WP2b覆核嘅新問題，令test fail。"""
    import csv
    from importer.wp3_run import run_wp3
    run_wp3()
    with open(paths.WP3_PAYROLL_DIFF_CSV, encoding="utf-8") as fh:
        rows = list(csv.DictReader(fh))
    page_mismatches = [r for r in rows if r["site_label"] == "__PAGE_TOTAL__"]
    for r in page_mismatches:
        assert r["category"] == "shift_count_diff", f"page total mismatch冇岩嘅category：{r}"
        policy_total = abs(Decimal(r["payroll_amount"]))
        diff = abs(Decimal(r["diff"]))
        assert policy_total > 0 and diff / policy_total < Decimal("0.05"), \
            f"逐頁總更數殘差超過policy總更數5%，可能係新問題唔係已知細差：{r}"


@pytestmark_e2e
def test_payroll_diff_has_no_tool_bug():
    """B目標：payroll-diff.csv 嘅tool_bug類別要係0（近似名/mixed_row/shift_
    count_diff/ot_or_adjustment都已經有專門category解釋，剩返嘅先算真正
    tool層面嘅未解釋差異）。"""
    import csv
    from importer.wp3_run import run_wp3
    run_wp3()
    with open(paths.WP3_PAYROLL_DIFF_CSV, encoding="utf-8") as fh:
        rows = list(csv.DictReader(fh))
    tool_bugs = [r for r in rows if r["category"] == "tool_bug"]
    assert not tool_bugs, f"payroll-diff.csv 有未分類嘅tool_bug：{tool_bugs}"


# ---------------------------------------------------------------------------
# WP5g（P0-執行計劃.md §8 Q10 已答）：master_roster_payout（A4）出街價最後
# fallback——總表逐崗位表頭外判價（out/wp5/master_posts.csv payout_rate），
# 糧明細/master_roster_rate/master_roster_only_rate 都搵唔到 guard_wage/
# broker_fee 先至用。之前 5 個長工（HKU Y更 G0061、數碼港CP1-4 G0139、
# 數碼港CP5 G0112、啟航1331 G0047/G0145，PDPO：只記guard_id，具體人可查
# out/wp2b/guards.csv）合共130更 guard_wage/broker_fee 空白（client_rate
# 有），rate_cards.csv S006 HKU Y更完全冇出街價卡。
# ---------------------------------------------------------------------------

@pytestmark_e2e
def test_master_roster_payout_fills_named_longterm_guard_blanks():
    """Regression test：呢5位安員（見上面模組註解嘅guard_id）之前因為
    master_roster_rate（A2）淨係覆蓋咗啟航1331K桶、S006 HKU Y更完全冇
    出街價卡，導致guard_wage_snapshot/broker_fee_snapshot雙雙空白（雖然
    client_rate已經有）。WP5g加咗master_roster_payout最後fallback之後：
      - G0061（S006 HKU Y更，桶K，26更）：guard_wage_snapshot應該全部係
        `850.00`（總表逐崗位表頭外判價，見rate_cards.csv source=
        master_roster_payout 嘅「總表崗位卡」）。
      - G0047／G0145（啟航1331 早更保安員，桶Y公司兵，各26更）：應該攞到
        總表釘死俾呢個guard_id嘅「總表長工卡」（guard_wage_source=
        monthly_salary_daily_avg），唔會撞用第二個人嘅價（兩人金額唔同：
        840.38 vs 718.02，如果撞埋一齊用會即刻喺呢度露餡）。
      - G0139（數碼港CP1-4，桶Y公司兵，26更）：總表崗位卡`730.77`。
    呢4位一共104更之前全部空白，而家應該0空白；broker_fee_snapshot 應該
    全部係`0.00`（總表外判價已經全包，唔再拆雞頭費）。"""
    import csv
    from importer.wp3_run import run_wp3
    run_wp3()
    with open(paths.WP3_SHIFT_RECORDS_PRICED_CSV, encoding="utf-8") as fh:
        rows = list(csv.DictReader(fh))

    expect = {
        "G0061": ("S006", Decimal("850.00"), 26),
        "G0047": ("S017", Decimal("840.38"), 26),
        "G0145": ("S017", Decimal("718.02"), 26),
        "G0139": ("S010", Decimal("730.77"), 26),
    }
    for guard_id, (site_id, expected_wage, expected_n) in expect.items():
        sub = [r for r in rows if r["guard_id"] == guard_id and r["site_id"] == site_id
               and r["record_type"] == "shift"]
        assert len(sub) == expected_n, f"{guard_id}@{site_id} 應該有{expected_n}更，實際{len(sub)}"
        blank = [r for r in sub if not r["guard_wage_snapshot"]]
        assert not blank, f"{guard_id}@{site_id} 仲有guard_wage_snapshot空白：{blank}"
        wages = {Decimal(r["guard_wage_snapshot"]) for r in sub}
        assert wages == {expected_wage}, f"{guard_id}@{site_id} guard_wage_snapshot應該全部係{expected_wage}，實際{wages}"
        fees = {r["broker_fee_snapshot"] for r in sub}
        assert fees == {"0.00"}, f"{guard_id}@{site_id} broker_fee_snapshot應該全部係0.00（總表外判價全包），實際{fees}"

    # rate_cards.csv 要有 master_roster_payout 嚟源嘅卡（S006 HKU Y更之前
    # 完全冇任何出街價卡——見P0-WP5g任務背景）。
    with open(paths.WP3_RATE_CARDS_CSV, encoding="utf-8") as fh:
        rate_cards = list(csv.DictReader(fh))
    mrp_cards = [rc for rc in rate_cards if rc["source"] == "master_roster_payout"]
    assert mrp_cards, "rate_cards.csv 應該有 source=master_roster_payout 嘅卡"
    s006_cards = [rc for rc in mrp_cards if rc["site_id"] == "S006"]
    assert s006_cards, "S006（HKU Y更）之前完全冇出街價卡，WP5g應該生成至少一張"


@pytestmark_e2e
def test_monthly_diff_known_reasons_are_populated():
    """P0-執行計劃WP3c §C：monthly-diff.csv 逐盤已知差異（category=="known"）
    全部要有非空嘅reason（唔可以淨係標known但冇解釋文字）；差異但冇已知
    原因嘅先至准標needs_review。

    WP5e應修7（2026-09-10，Opus最終複驗）：加埋needs_review上限——目前
    已知淨低1條（S026／創富廠IC，shift_records 22 vs R2 24，缺口2更，
    「資料流缺口，交WP4」，WP4層面已經喺`_L1_SITE_OVERRIDE["S026"]`接
    手覆核，唔係查唔到出處）。呢條上限刻意寫死做1，係一個防回歸嘅
    高水位警報：如果WP1/WP2b/WP3邏輯改動令needs_review數量*增加*（多咗
    未覆核嘅缺口），呢個test應該即刻爆，逼人跟返WP4個「零未解釋」收貨
    標準去追加覆核，唔可以靜靜由needs_review條數飆升。"""
    import csv
    from importer.wp3_run import run_wp3
    run_wp3()
    with open(paths.WP3_MONTHLY_DIFF_CSV, encoding="utf-8") as fh:
        rows = list(csv.DictReader(fh))
    assert rows, "monthly-diff.csv 唔應該係空"
    known_no_reason = [r for r in rows if r["category"] == "known" and not r.get("reason", "").strip()]
    assert not known_no_reason, f"category=known 但冇reason嘅行：{known_no_reason}"
    bad_categories = {r["category"] for r in rows} - {"known", "needs_review"}
    assert not bad_categories, f"monthly-diff.csv 出現預期外嘅category值：{bad_categories}"
    needs_review = [r for r in rows if r["category"] == "needs_review"]
    _KNOWN_NEEDS_REVIEW_SITE_IDS = {"S026"}  # 創富廠IC，WP4已接手覆核（見_L1_SITE_OVERRIDE）
    assert len(needs_review) <= 1, (
        f"needs_review 上限係1條（目前已知淨低{sorted(_KNOWN_NEEDS_REVIEW_SITE_IDS)}），"
        f"而家有{len(needs_review)}條，多咗嘅要人手追加覆核先可以加大呢個上限：{needs_review}"
    )
    unexpected_needs_review = {r["site_id"] for r in needs_review} - _KNOWN_NEEDS_REVIEW_SITE_IDS
    assert not unexpected_needs_review, (
        f"出現咗未預期嘅needs_review site_id（唔係已知嘅{sorted(_KNOWN_NEEDS_REVIEW_SITE_IDS)}），"
        f"要人手覆核先可以更新呢個清單：{unexpected_needs_review}"
    )


@pytestmark_e2e
def test_run_wp3_three_times_stable_xlsx_content():
    """xlsx內部zip entry嘅timestamp每次save都唔同（openpyxl行為，_save_
    stable淨係定死docProps嗰兩個timestamp，唔控制zip層），所以raw md5唔
    可以直接比較（同`test_run_wp3_three_times_byte_identical_csv`嗰個csv
    md5唔同）——呢度用內容hash（逐sheet／逐cell嘅值），對應P0-執行計劃
    WP3c品質要求「三次重跑…xlsx用內容hash」。"""
    import hashlib
    from openpyxl import load_workbook
    from importer.wp3_run import run_wp3

    xlsx_paths = [
        paths.OUT_WP3_REPORTS_DIR / "月結總表-regenerated.xlsx",
        paths.OUT_WP3_REPORTS_DIR / "payroll" / "雞頭糧明細-regenerated.xlsx",
        paths.OUT_WP3_REPORTS_DIR / "roster" / _STANDARD_FILE,
        paths.OUT_WP3_REPORTS_DIR / "roster" / _CP_FILE,
        paths.OUT_WP3_REPORTS_DIR / "roster" / _CITYBUS_FILE,
    ]

    def _content_hash():
        h = hashlib.md5()
        for p in xlsx_paths:
            wb = load_workbook(p, data_only=True)
            for ws in wb.worksheets:
                h.update(ws.title.encode("utf-8"))
                for row in ws.iter_rows():
                    for cell in row:
                        h.update(repr(cell.value).encode("utf-8"))
        return h.hexdigest()

    run_wp3()
    h1 = _content_hash()
    run_wp3()
    h2 = _content_hash()
    run_wp3()
    h3 = _content_hash()
    assert h1 == h2 == h3
