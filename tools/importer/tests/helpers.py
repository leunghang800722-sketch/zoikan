"""Small helpers for building synthetic worksheets in tests, so layout /
classification logic can be tested without touching the real (gitignored)
data/2026-07 xlsx files.
"""
import datetime

import openpyxl


def july_date(day: int) -> datetime.datetime:
    return datetime.datetime(2026, 7, day)


def new_sheet():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "2026"
    return wb, ws


def write_row(ws, row: int, values: dict):
    """values: {col_index(1-based): value}"""
    for col, val in values.items():
        ws.cell(row=row, column=col, value=val)


def write_date_row(ws, row: int, label_col: int, n_days: int = 31, start_day: int = 1):
    ws.cell(row=row, column=label_col, value="Date")
    for i in range(n_days):
        ws.cell(row=row, column=label_col + 1 + i, value=july_date(start_day + i))
