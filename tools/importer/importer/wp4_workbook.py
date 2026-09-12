"""WP4 老闆工作簿 + technical 工作簿產生器。

老闆版 `P0-對數工作簿-2026-07.xlsx`：全部中文標籤，**唔准出現
site_id／guard_id／rate_id 呢類內部代碼，亦唔准出現 `.csv` 檔名／
`gate=`／`mixed_site_label`／`monthly_fee_day`／`day-mismatch` 呢類工程
術語**（`_boss_text` 負責轉換／剝走）；`工號`（Uxxxxx 呢種員工工號）唔
算內部代碼，係老闆本身糧單/更表已經用緊嘅業務代號，保留。

Technical 版 `P0-對數工作簿-2026-07-technical.xlsx`：內容一樣但唔轉換／
唔剝code，俾要追溯出處嘅人（開發／Opus 驗收）用。
"""
import re
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import List

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from . import paths

# openpyxl 預設將「而家」時間寫入 docProps/core.xml（created/modified），
# 令成個 xlsx 每次 save 都 byte 唔同，就算內容一模一樣。跟 WP3
# `wp3_reports._save_stable` 同一原則（WP4 唔可以改WP3源code，呢度自己
# 喺wp4_workbook.py複製一份），save之前將呢兩個欄位定死做常數。
_FIXED_TIMESTAMP = datetime(2026, 7, 31, 0, 0, 0)


def _save_stable(wb: Workbook, path: Path) -> None:
    wb.properties.created = _FIXED_TIMESTAMP
    wb.properties.modified = _FIXED_TIMESTAMP
    wb.properties.creator = "importer(wp4)"
    wb.properties.lastModifiedBy = "importer(wp4)"
    wb.save(path)
    # 2026-09-10 發現：openpyxl 嘅 save_workbook()（writer/excel.py）save 之
    # 前會強制 `workbook.properties.modified = datetime.now()`，完全蓋過
    # 上面設低嘅 _FIXED_TIMESTAMP——即係話上面四行對 docProps/core.xml 嘅
    # modified 欄完全冇用，每次 save 個 xlsx 都會因為呢一個內部 timestamp
    # byte 唔同（連跑3次md5理應一樣呢個驗收條件實際做唔到）。呢個係
    # openpyxl library 本身嘅行為，唔關內容/數據事——save完之後直接重寫
    # zip 入面嗰份 docProps/core.xml，逼佢用返 _FIXED_TIMESTAMP，先至真係
    # byte-for-byte 穩定。
    _force_stable_core_xml(path)


def _force_stable_core_xml(path: Path) -> None:
    import zipfile
    with zipfile.ZipFile(path, "r") as zin:
        infos = zin.infolist()
        contents = {info.filename: zin.read(info.filename) for info in infos}
    core_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
        'xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/" '
        'xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
        '<dc:creator>importer(wp4)</dc:creator>'
        f'<dcterms:created xsi:type="dcterms:W3CDTF">{_FIXED_TIMESTAMP.isoformat()}Z</dcterms:created>'
        f'<dcterms:modified xsi:type="dcterms:W3CDTF">{_FIXED_TIMESTAMP.isoformat()}Z</dcterms:modified>'
        '<cp:lastModifiedBy>importer(wp4)</cp:lastModifiedBy>'
        '</cp:coreProperties>'
    ).encode("utf-8")
    contents["docProps/core.xml"] = core_xml
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in sorted(infos, key=lambda i: i.filename):
            new_info = zipfile.ZipInfo(info.filename, date_time=(2026, 7, 31, 0, 0, 0))
            new_info.compress_type = zipfile.ZIP_DEFLATED
            new_info.external_attr = info.external_attr
            zout.writestr(new_info, contents[info.filename])

_HEADER_FILL = PatternFill(start_color="FFDDEBF7", end_color="FFDDEBF7", fill_type="solid")
_HEADER_FONT = Font(bold=True)
_WRAP = Alignment(wrap_text=True, vertical="top")

_STATUS_ZH = {"explained": "已解釋", "awaiting_boss": "待老闆決定"}
_LAYER_ZH = {"L1": "L1 更數", "L2": "L2 成本", "L3": "L3 收入"}

# ---------------------------------------------------------------------------
# 老闆版文字轉換（技術詞 → 人話中文）
# ---------------------------------------------------------------------------
# 呢個轉換表要喺「剝碼」（下面嘅 _CODE_*_RE）之前套用——好多技術詞本身
# 就係整句嘅一部分（例如「已摺埋入S006嘅billing_units.csv
# extra_shifts_from_monthly」呢句嚟自 wp3_reports.py，WP4 唔准改WP3源
# code），如果淨係剝走個 site code 就會剩低「已摺埋入嘅billing_units.csv
# extra_shifts_from_monthly」呢種斷咗頭嘅爛句；要先將成個技術詞組轉做
# 完整中文詞，先至剝碼，句子先讀得通。
_TERM_MAP = [
    # 修必修5（2026-09-10，WP5c）：擴充呢個表，剝走 WP5-C 之後新出現嘅
    # 技術詞／內部代號／PDPO 註腳／out/data 路徑——全部要喺「剝碼」
    # （_CODE_TOKEN_RE 等）之前處理，順序見下面逐條。
    #
    # PDPO 註腳（「（PDPO：真名唔喺source code，可查out/wp2b/guards.csv）」
    # 呢類括號註解）——成句連埋佢引用嘅路徑一齊剝走，唔淨低斷咗頭嘅逗號。
    (re.compile(r"[，,；]?\s*PDPO：[^）；]*"), ""),
    (re.compile(r"PDPO"), "個人資料"),
    # WP5e必修3（2026-09-10，Opus最終複驗）：老闆版清乾淨剩低嘅一批技術詞
    # ／內部代號——順序原則同上面一致：越specific（成句/成個複合詞）嘅
    # entry要行喺越generic嘅之前，避免generic先食走部分字元、剩低斷咗頭
    # 嘅殘留（例如「autopay_longterm」一定要喺bare「autopay」之前處理，
    # 唔係就會剩低半桶水嘅「自動轉賬_longterm」）。
    (re.compile(r"out_of_scope_or_unmapped"), "範圍外或未配對"),
    (re.compile(r"autopay_longterm"), "自動轉賬長工"),
    (re.compile(r"is_longterm=True"), "屬長工"),
    (re.compile(r"is_longterm=False"), "唔屬長工"),
    (re.compile(r"(?<![A-Za-z0-9_])is_longterm(?![A-Za-z0-9_])"), "是咪長工"),
    (re.compile(r"(?<![A-Za-z0-9_])settlement_unit(?![A-Za-z0-9_])"), "結算單位"),
    (re.compile(r"shift_type_id"), "更別編號"),
    # WP5e必修3注意：「Σshifts where shift_kind=normal」呢句成句轉換
    # （原本喺呢個檔案下面）一定要行喺bare「shift_kind」轉換之前，唔係
    # bare嗰條會搶先將"shift_kind"轉走，令成句pattern永遠match唔到，剩
    # 低斷咗頭嘅「Σshifts where 更次種類=normal」。
    (re.compile(r"Σshifts where shift_kind=normal"), "口徑更數計法（剔除OT/調整行嘅正更加總）"),
    (re.compile(r"(?<![A-Za-z0-9_])shift_kind(?![A-Za-z0-9_])"), "更次種類"),
    (re.compile(r"(?<![A-Za-z0-9_])payroll_ref(?![A-Za-z0-9_])"), "糧單出處"),
    (re.compile(r"(?<![A-Za-z0-9_])client_rate(?![A-Za-z0-9_])"), "收客價"),
    (re.compile(r"(?<![A-Za-z0-9_])payout_rate(?![A-Za-z0-9_])"), "出街價"),
    (re.compile(r"(?<![A-Za-z0-9_])known_shift_gap(?![A-Za-z0-9_])"), "已知更數差"),
    (re.compile(r"(?<![A-Za-z0-9_])quantize[ds]?(?![A-Za-z0-9_])"), "尾數處理"),
    (re.compile(r"(?<![A-Za-z0-9_])resolved(?![A-Za-z0-9_])"), "已配對"),
    (re.compile(r"(?<![A-Za-z0-9_])resolve(?![A-Za-z0-9_])"), "配對"),
    (re.compile(r"(?<![A-Za-z0-9_])regenerated(?![A-Za-z0-9_])"), "重新生成"),
    (re.compile(r"(?<![A-Za-z0-9_])regen_revenue(?![A-Za-z0-9_])"), "我哋計嘅收入"),
    (re.compile(r"(?<![A-Za-z0-9_])regen(?![A-Za-z0-9_])"), "系統重新計嘅"),
    (re.compile(r"Y-coded"), "標Y代號嘅"),
    (re.compile(r"build_l[123]_[a-z_]+"), "系統內部分類邏輯"),
    (re.compile(r"(?<![A-Za-z0-9_])row_no(?![A-Za-z0-9_])"), "行號"),
    (re.compile(r"(?<![A-Za-z0-9_])master_roster_only(?![A-Za-z0-9_])"), "只喺總表出現"),
    (re.compile(r"(?<![A-Za-z0-9_])master_roster_rate(?![A-Za-z0-9_])"), "總表價目"),
    (re.compile(r"(?<![A-Za-z0-9_])master_roster[a-z_]*(?![A-Za-z0-9_])"), "總表"),
    (re.compile(r"(?<![A-Za-z0-9_])client_group(?![A-Za-z0-9_])"), "客戶集團"),
    (re.compile(r"(?<![A-Za-z0-9_])cell_kind(?![A-Za-z0-9_])"), "格仔類型"),
    (re.compile(r"(?<![A-Za-z0-9_])site_ids(?![A-Za-z0-9_])"), "盤代號清單"),
    (re.compile(r"(?<![A-Za-z0-9_])payroll_only(?![A-Za-z0-9_])"), "只喺糧單出現"),
    (re.compile(r"(?<![A-Za-z0-9_])flat_referral(?![A-Za-z0-9_])"), "固定介紹費"),
    (re.compile(r"(?<![A-Za-z0-9_])near_name(?![A-Za-z0-9_])"), "近似姓名個案"),
    (re.compile(r"(?<![A-Za-z0-9_])roster_name(?![A-Za-z0-9_])"), "更表姓名"),
    (re.compile(r"(?<![A-Za-z0-9_])per_shift(?![A-Za-z0-9_])"), "逐更"),
    (re.compile(r"(?<![A-Za-z0-9_])wp3_run(?![A-Za-z0-9_])"), "系統定價邏輯"),
    (re.compile(r"=True(?![A-Za-z0-9_])"), "：是"),
    (re.compile(r"=False(?![A-Za-z0-9_])"), "：否"),
    (re.compile(r"(?<![A-Za-z0-9_])True(?![A-Za-z0-9_])"), "是"),
    (re.compile(r"(?<![A-Za-z0-9_])False(?![A-Za-z0-9_])"), "否"),
    # rule 0.1／0.2（WP5-C settlement_page 修）：轉做人話。要喺「WP\d」
    # 呢條通用規則之前處理（「rule」後面跟緊數字，唔會撞落WP\d嗰條）。
    (re.compile(r"rule\s*0\.1"), "總表逐日安員歸屬確認"),
    (re.compile(r"rule\s*0\.2"), "總表逐日盤點配額確認"),
    (re.compile(r"rule\s*1\.55"), "長工月薪表配對"),
    (re.compile(r"rule\s*1\.6"), "盤級糧單規則"),
    (re.compile(r"rule\s*\d(?:\.\d+)?"), "配對規則"),
    # WP5-C／WP2b-c 呢類「WP + 數字 + 字母 + 連字號+字母」複合代號：舊
    # regex `WP\d[abc]?` 淨食到「WP5」剩低「-C」冇轉換，變成半唐番嘅
    # 「工作包-C」；擴闊到食埋trailing嘅連字號+字母／字母尾綴。
    (re.compile(r"WP\d[a-zA-Z]?(?:-[A-Za-z])?"), "工作包"),
    (re.compile(r"edit distance"), "姓名相似度"),
    (re.compile(r"site override"), "跨盤指派"),
    (re.compile(r"block\s*標題"), "分組標題"),
    (re.compile(r"billing_model"), "計費模式"),
    (re.compile(r"settlement_page"), "糧單頁面"),
    (re.compile(r"master_broker_code"), "總表雞頭代號"),
    (re.compile(r"master_broker_source"), "總表判斷來源"),
    (re.compile(r"broker_id"), "雞頭編號"),
    (re.compile(r"parent_shift_id"), "母更次編號"),
    (re.compile(r"ot_hours"), "OT時數"),
    (re.compile(r"Σshifts where shift_kind=normal"), "口徑更數計法（剔除OT/調整行嘅正更加總）"),
    (re.compile(r"unresolved_guards"), "未確認結算單位安員名單"),
    (re.compile(r"policy_total_shifts"), "口徑更數"),
    (re.compile(r"rate_cards?"), "價目表"),
    (re.compile(r"site_label"), "盤名標籤"),
    (re.compile(r"billing_units?"), "開票單位表"),
    (re.compile(r"shift_type"), "更別"),
    (re.compile(r"special"), "特殊收入項"),
    (re.compile(r"base_[a-z_]+"), "基準數"),
    (re.compile(r"payroll-diff"), "糧單逐行核對"),
    (re.compile(r"monthly-diff"), "月結逐盤核對"),
    (re.compile(r"roundtrip-diff"), "更表逐格核對"),
    # 修必修3（2026-09-10）：舊regex用`\b`——Python嘅`\w`／`\b`當中文字係
    # word char，「喺row21」呢種中文字直接貼住「row」嘅情況，中文字同
    # 「r」之間根本冇word boundary，`\brow(\d+)`會匹配失敗、漏走個
    # 「row21」。改用ASCII-only嘅前後負向lookaround（同呢個檔案其他
    # identifier一致嘅做法）。
    (re.compile(r"(?<![A-Za-z0-9_])row(\d+)(?![A-Za-z0-9_])"), r"第\1行"),
    (re.compile(r"(?<![A-Za-z0-9_])cost="), "成本"),
    (re.compile(r"in-scope"), "範圍內"),
    (re.compile(r"out_of_scope"), "範圍外"),
    (re.compile(r"explained"), "已解釋"),
    (re.compile(r"awaiting_boss"), "待老闆"),
    (re.compile(r"Σ\(guard_wage_snapshot\+broker_fee_snapshot\)"), "我哋計出嘅成本（工資＋雞頭費）"),
    (re.compile(r"guard_wage"), "工資"),
    (re.compile(r"broker_fee"), "雞頭費"),
    (re.compile(r"shift_records\.csv"), "系統更次記錄"),
    (re.compile(r"shift_records"), "系統更次記錄"),
    (re.compile(r"snapshot"), "系統計算"),
    (re.compile(r"staff_no"), "工號"),
    (re.compile(r"unknown"), "未知"),
    # 修必修5（2026-09-10）：舊regex `guard(?!_)` 淨係防止匹配到
    # `guard_xxx`（後面跟底線），但冇防前面／後面貼住其他英文字母/數字/
    # 底線嘅情況——「unresolved_guards」入面嘅「guard」後面跟緊「s」（唔
    # 係底線），會被錯誤命中，變成半唐番嘅「unresolved_安員s」。改用完整
    # 前後負向lookaround，淨係匹配獨立一個word嘅「guard」。
    (re.compile(r"(?<![A-Za-z0-9_])guard(?![A-Za-z0-9_])"), "安員"),
    (re.compile(r"(?<![A-Za-z0-9_])guards(?![A-Za-z0-9_])"), "安員名單"),
    (re.compile(r"unresolved_guards"), "未確認結算單位安員名單"),
    (re.compile(r"needs_review"), "未分類"),
    (re.compile(r"autopay"), "自動轉賬"),
    (re.compile(r"\bR2\b"), "月結R2"),
    # 呢句成句一齊處理（wp3_reports.py 生成，WP4 唔可以改WP3源code，
    # 淨係喺呢層做轉換）：「已摺埋入S006嘅billing_units.csv
    # extra_shifts_from_monthly，唔係獨立site」→ 完整人話句。
    (re.compile(r"已摺埋入(?:S\d{3})?嘅billing_units\.csv extra_shifts_from_monthly，唔係獨立(?:site|盤)"),
     "已摺入主結算單位嘅月結副行，唔係獨立盤"),
    # 另一句成句轉換（都係 wp3_reports.py 生成、payroll-diff.csv 嘅
    # mixed_site_label note，一樣唔可以改WP3源code）：「橫跨
    # site_id=['S005', 'S019']」呢種 Python list literal 唔係人話，成句
    # 轉做「橫跨幾個唔同盤」——一定要行喺 code-token 剝除之前，唔係就會
    # 剩低「橫跨site_id=['', '']」呢種斷咗頭嘅殘留。
    (re.compile(r"橫跨site_id=\[[^\]]*\]"), "橫跨幾個唔同盤"),
    (re.compile(r"每個site_id"), "每個盤"),
    (re.compile(r"billing_units\.csv"), "月結副行對照表"),
    (re.compile(r"extra_shifts_from_monthly"), "月結副行更次"),
    (re.compile(r"shift_records_priced\.csv"), "已計價更表記錄"),
    (re.compile(r"shift_records\.csv"), "更表系統記錄"),
    (re.compile(r"day-mismatch\.csv"), "更表／糧單逐日核對記錄"),
    (re.compile(r"day-mismatch"), "更表／糧單逐日核對"),
    (re.compile(r"payroll-diff\.csv"), "糧單對數記錄"),
    (re.compile(r"payroll-pages\.csv"), "糧單頁面記錄"),
    (re.compile(r"monthly-diff\.csv"), "月結對數記錄"),
    (re.compile(r"gate-results\.csv"), "開票查核記錄"),
    (re.compile(r"billing-unit-solve\.csv"), "開票單位價目記錄"),
    (re.compile(r"billing_adjustments\.csv"), "細殘差記錄"),
    (re.compile(r"name-near-match\.csv"), "近似姓名配對記錄"),
    (re.compile(r"unresolved-guards\.csv"), "未確認安員記錄"),
    (re.compile(r"kai1331-daily-residuals\.csv"), "啟航1331逐日核算記錄"),
    (re.compile(r"broker-conflicts\.csv"), "雞頭歸屬衝突記錄"),
    (re.compile(r"guards\.csv"), "安員名單記錄"),
    (re.compile(r"adjustments\.csv"), "調整記錄"),
    (re.compile(r"validate-identity-conflicts\.csv"), "身份衝突記錄"),
    (re.compile(r"mixed_site_label"), "跨盤合併標籤"),
    (re.compile(r"monthly_fee_day"), "月費制人日"),
    (re.compile(r"(?<![A-Za-z0-9_])monthly_fee(?![A-Za-z0-9_])"), "月費制"),
    (re.compile(r"monthly_fixed"), "月費制"),
    (re.compile(r"admin_fee"), "行政費制"),
    (re.compile(r"gate=PASS\([^)]*\)"), "查核已通過"),
    (re.compile(r"gate="), "查核狀態："),
    # 注意：呢度刻意唔用 `\b`——Python 嘅 `\w`／`\b` 對中文字係當「word
    # char」，`site_or_unit` 呢類 identifier 好多時直接貼住中文字（例如
    # 「見site_or_unit相同」），`見`同`s`之間根本冇 word boundary，
    # `\bsite_or_unit\b` 會匹配失敗、留低個英文identifier漏出嚟。呢啲
    # identifier 本身都夠獨特（唔會係第二個字嘅子字串），淨係用
    # ASCII-only嘅前後負向lookaround做最基本防護就夠。
    (re.compile(r"(?<![A-Za-z0-9_])site_or_unit(?![A-Za-z0-9_])"), "盤／單位"),
    (re.compile(r"(?<![A-Za-z0-9_])guard_id(?![A-Za-z0-9_])"), "呢個安員"),
    (re.compile(r"(?<![A-Za-z0-9_])site_id(?![A-Za-z0-9_])"), "呢個盤"),
    (re.compile(r"(?<![A-Za-z0-9_])rate_id(?![A-Za-z0-9_])"), "呢張價目卡"),
]
# 冇喺上面表入面、但仍然係內部檔名嘅殘留（.csv/.xlsx/.xls/.pdf/.md），
# 一律隱形剝走（唔留低半句英文檔名俾老闆睇）。刻意唔用 `\b`（同上，中文
# 字前後冇word boundary），stem 淨係揀ASCII/連字號/底線（唔撈中文），
# 咁樣先唔會貪心咬埋前面嘅中文字（filename本身喺呢個project一律純英文
# 加數字）。
_GENERIC_FILE_RE = re.compile(r"[A-Za-z0-9_\-]+\.(?:csv|xlsx|xls|pdf|md|py)")
# 修必修3（2026-09-10）：中文化咗嘅檔名（例如「安員名單.csv」由
# `guards.csv`轉換嚟）唔會撞落上面嗰條（stem要求ASCII），淨低個
# 「.csv」尾綴冇剝走。呢度淨剝副檔名本身（唔理前面個stem係中文定英文），
# 一律隱形剝走，留低個檔名主體（「安員名單.csv」→「安員名單」）。
_CN_FILE_EXT_RE = re.compile(r"\.(?:csv|xlsx|xls|pdf|md|py)(?![A-Za-z0-9_])")
# 修必修5（2026-09-10）：out/／data/ 開頭嘅目錄路徑（例如
# 「可查out/wp2b/guards.csv」剝走檔名之後剩低嘅「out/wp2b/」，或者根本
# 冇檔名淨係目錄，例如「data/mappings-private/」）——呢啲對老闆嚟講都係
# 冇意義嘅內部路徑，一律隱形剝走。要喺 `_GENERIC_FILE_RE` 之後行（等
# 檔名先剝走，避免呢條規則貪心咬走檔名入面嘅字元）。
_GENERIC_PATH_RE = re.compile(r"(?:out|data)/[A-Za-z0-9_\-./]*")

# 一串用「/」或「、」連埋嘅內部代碼（例如 G0157/G0033）要成組一齊剝走，
# 唔可以逐個token剝、留低斷咗頭嘅分隔符（「得G0157/G0033兩個人輪」剝到
# 「得/兩個人輪」）。呢個規則刻意淨係match code token 之間嘅分隔符，唔
# 會誤中「安員/日」呢種正常中文用嘅斜綫（因為兩邊都要係 code token 先
# match 到）。
_CODE_TOKEN = r"(?:SR\d{6}|ST\d{4}|RC\d{4}|G\d{4}|S\d{3}|B\d{2})"
_CODE_TOKEN_SEQ_RE = re.compile(_CODE_TOKEN + r"(?:[/、,，]" + _CODE_TOKEN + r")+")
_CODE_KV_RE = re.compile(r"(?<![A-Za-z0-9_])\w+_id=\S+?(?=[;；,，\s)）]|$)")
_CODE_TOKEN_RE = re.compile(r"(?<![A-Za-z0-9])" + _CODE_TOKEN + r"(?![A-Za-z0-9])")
# WP5e必修3（2026-09-10，Opus最終複驗）：防線最後一層——`_TERM_MAP`轉換
# 表淨係人手列咗識別到嘅技術詞，將來WP1-3上游報表可能會再生成新嘅
# snake_case內部欄位名（呢個模組唔可以改WP1-3源code），漏咗轉換就會
# 老闆版出現生嘅英文identifier。呢度加一條通用防線：轉換表行完之後，
# 淨低嘅ASCII snake_case token（兩截或以上、純細楷英文/數字加底線）一律
# 隱形剝走，等老闆版100%冇機會漏出「未經翻譯」嘅工程術語（同下面
# test_boss_workbook_has_no_technical_jargon()同一把尺）。真.工號
# 「Uxxxxx」呢類業務代號係大楷開頭+數字，唔會撞落呢條規則。
_SNAKE_CASE_FALLBACK_RE = re.compile(r"(?<![A-Za-z0-9])[a-z][a-z0-9]*(?:_[a-z0-9]+)+(?![A-Za-z0-9])")
_EMPTY_PAREN_RE = re.compile(r"[（(]\s*[）)]")
_DUP_PUNCT_RE = re.compile(r"[;；]{2,}")
_DUP_SPACE_RE = re.compile(r" {2,}")


def _boss_text(text) -> str:
    """老闆版文字轉換：先將技術詞組轉做人話中文（`_TERM_MAP`），先至剝
    走 site_id／guard_id／rate_id 呢類內部代碼；`工號 Uxxxxx` 呢種真實
    業務工號（老闆糧單/更表本身已經用緊）唔會被呢個 regex 命中，照留。
    剝碼順序刻意「成組代碼一齊剝」行先（`_CODE_TOKEN_SEQ_RE`），先至
    「單個代碼」（`_CODE_TOKEN_RE`），避免剝完淨低斷咗頭嘅分隔符
    （例如「G0157/G0033」剝出嚟唔會剩低「/」）。"""
    if text is None:
        return ""
    t = str(text)
    for pattern, repl in _TERM_MAP:
        t = pattern.sub(repl, t)
    t = _GENERIC_FILE_RE.sub("", t)
    t = _GENERIC_PATH_RE.sub("", t)
    t = _CN_FILE_EXT_RE.sub("", t)
    t = _CODE_TOKEN_SEQ_RE.sub("", t)
    t = _CODE_KV_RE.sub("", t)
    t = _CODE_TOKEN_RE.sub("", t)
    t = _SNAKE_CASE_FALLBACK_RE.sub("", t)
    t = _EMPTY_PAREN_RE.sub("", t)
    t = _DUP_PUNCT_RE.sub("；", t)
    t = _DUP_SPACE_RE.sub(" ", t)
    t = t.replace("（；", "（").replace("；）", "）")
    return t.strip(" ；;")


# 舊名保留做別名，以防其他模組／舊 import 仲用緊呢個名。
_scrub_boss_text = _boss_text


def _try_num_boss(v):
    n = _try_num(v)
    return _boss_text(n) if isinstance(n, str) else n


def _try_num(v):
    if v is None or v == "":
        return None
    try:
        dv = Decimal(str(v).replace(",", ""))
        return float(dv) if "." in str(v) or dv != dv.to_integral_value() else int(dv)
    except (InvalidOperation, ValueError):
        return v


def _autosize(ws, widths: List[int]):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _header_row(ws, headers: List[str]):
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
    ws.freeze_panes = "A2"


def _summary_rows(summary: dict):
    lines = []
    for layer in ("L1", "L2", "L3"):
        lc = summary["by_layer"].get(layer, {})
        explained = lc.get("explained", 0)
        awaiting = lc.get("awaiting_boss", 0)
        lines.append((_LAYER_ZH[layer], explained, awaiting, explained + awaiting))
    return lines


def build_workbooks(summary: dict, diff_rows: List[dict], d: dict, boss_facts: dict):
    _build_one(summary, diff_rows, d, boss_facts, boss=True, path=paths.WP4_WORKBOOK_XLSX)
    _build_one(summary, diff_rows, d, boss_facts, boss=False, path=paths.WP4_WORKBOOK_TECHNICAL_XLSX)


def _build_one(summary: dict, diff_rows: List[dict], d: dict, boss_facts: dict, boss: bool, path: Path):
    wb = Workbook()

    # ---- Sheet 1：總覽 ----
    ws = wb.active
    ws.title = "總覽"
    ws.append(["P0 三層對數總覽 — 2026年7月"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(["層", "已解釋（項）", "待老闆決定（項）", "合計（項）"])
    for c in range(1, 5):
        ws.cell(row=3, column=c).font = _HEADER_FONT
        ws.cell(row=3, column=c).fill = _HEADER_FILL
    for row in _summary_rows(summary):
        ws.append(list(row))
    ws.append([])
    ws.append(["收貨標準：三層差異清單入面零「未解釋」項目（唔可以有查完都答唔到「呢條係乜」嘅項）。"])
    ws.append([f"差異清單總項數：{len(diff_rows)}；當中零項標「需要覆核未分類」，"
               f"每一項都已經分類 + 有出處，要麼「已解釋」，要麼「待老闆決定」。"])
    ws.append([])
    ws.append(["分類", "項數"])
    ws.cell(row=ws.max_row, column=1).font = _HEADER_FONT
    ws.cell(row=ws.max_row, column=2).font = _HEADER_FONT
    for cat, n in sorted(summary["by_category"].items(), key=lambda kv: -kv[1]):
        ws.append([cat, n])
    ws.append([])
    autopay_row = next((r for r in diff_rows if r["site_or_unit"] == "長工 autopay（BOCHK D2）"), None)
    if autopay_row is not None:
        ws.append([f"提示：「長工 autopay（BOCHK D2）」呢行差額（約{autopay_row['diff']}元）係名義攤分"
                    "（用月結出街價中位數／加權平均反推），唔係「數唔平」——長工出固定月薪（autopay），"
                    "唔係逐更計錢，呢個差額唔應該當做總覽嘅差異金額計落總數。"])
    _autosize(ws, [46, 16, 18, 12])

    # ---- Sheet 2-4：L1/L2/L3 ----
    for layer in ("L1", "L2", "L3"):
        ws = wb.create_sheet(_LAYER_ZH[layer])
        headers = (["盤／單位", "我哋（更表/系統）", "老闆嘅版本", "差額", "單位",
                     "分類", "說明", "建議處理", "狀態", "老闆問題編號"] if boss else
                    ["diff_id", "盤／單位", "我哋（更表/系統）", "老闆嘅版本", "差額", "單位",
                     "出處", "分類", "說明", "建議處理", "狀態", "老闆問題編號"])
        _header_row(ws, headers)
        for r in diff_rows:
            if r["layer"] != layer:
                continue
            status_zh = _STATUS_ZH.get(r["status"], r["status"])
            if boss:
                ws.append([
                    _boss_text(r["site_or_unit"]), _try_num_boss(r["ours"]), _try_num_boss(r["theirs"]),
                    _try_num(r["diff"]), r["unit"], r["category"], _boss_text(r["explanation"]),
                    _boss_text(r["recommended_action"]), status_zh, r["boss_question_ref"],
                ])
            else:
                ws.append([
                    r["diff_id"], r["site_or_unit"], _try_num(r["ours"]), _try_num(r["theirs"]),
                    _try_num(r["diff"]), r["unit"], r["provenance"], r["category"], r["explanation"],
                    r["recommended_action"], status_zh, r["boss_question_ref"],
                ])
            ws.cell(row=ws.max_row, column=(7 if boss else 9)).alignment = _WRAP
        widths = ([26, 14, 14, 12, 10, 30, 46, 30, 12, 10] if boss else
                  [8, 22, 14, 14, 12, 10, 40, 30, 46, 30, 12, 10])
        _autosize(ws, widths)

    # ---- Sheet 5：差異清單（全部，逐條）----
    # 老闆版明文剝走 diff_id（D0開頭代碼），technical版保留（俾開發/Opus
    # 驗收追溯出處用）。
    ws = wb.create_sheet("差異清單")
    headers = (["層", "盤／單位", "我哋", "老闆嘅版本", "差額", "單位", "分類",
                 "說明", "建議處理", "狀態", "老闆問題編號"] if boss else
                ["diff_id", "layer", "site_or_unit", "ours", "theirs", "diff", "unit",
                 "provenance", "category", "explanation", "recommended_action", "status", "boss_question_ref"])
    _header_row(ws, headers)
    for r in diff_rows:
        status_zh = _STATUS_ZH.get(r["status"], r["status"])
        if boss:
            ws.append([_LAYER_ZH[r["layer"]], _boss_text(r["site_or_unit"]),
                       _try_num_boss(r["ours"]), _try_num_boss(r["theirs"]), _try_num(r["diff"]), r["unit"],
                       r["category"], _boss_text(r["explanation"]), _boss_text(r["recommended_action"]),
                       status_zh, r["boss_question_ref"]])
        else:
            ws.append([r["diff_id"], r["layer"], r["site_or_unit"], _try_num(r["ours"]), _try_num(r["theirs"]),
                       _try_num(r["diff"]), r["unit"], r["provenance"], r["category"], r["explanation"],
                       r["recommended_action"], r["status"], r["boss_question_ref"]])
    widths = ([10, 26, 12, 12, 10, 10, 24, 42, 28, 12, 10] if boss else
              [8, 6, 22, 12, 12, 10, 10, 36, 24, 42, 28, 14, 10])
    _autosize(ws, widths)

    # ---- Sheet 6：待老闆決定（每行=一條問題，唔係逐條diff row；老闆版
    # 唔帶diff_id，technical版備註欄保留diff_id方便追溯）----
    ws = wb.create_sheet("待老闆決定")
    _header_row(ws, ["問題編號", "一句問題", "涉及安員或盤", "日期", "影響更數", "影響金額", "我哋建議"])
    for q in boss_facts["boss_questions"]:
        question, entities, recommendation = q["question"], q["entities"], q["recommendation"]
        dates, shift_impact, money_impact = q["dates"], q["shift_impact"], q["money_impact"]
        if boss:
            question = _boss_text(question)
            entities = _boss_text(entities)
            recommendation = _boss_text(recommendation)
            dates = _boss_text(dates)
            shift_impact = _boss_text(shift_impact)
            money_impact = _boss_text(money_impact)
        else:
            recommendation = f"{recommendation}（{q['n']}項；{q['diff_ids']}）"
        ws.append([q["ref"], question, entities, dates, shift_impact, money_impact, recommendation])
        ws.cell(row=ws.max_row, column=2).alignment = _WRAP
        ws.cell(row=ws.max_row, column=7).alignment = _WRAP
    _autosize(ws, [10, 52, 26, 22, 12, 14, 44])

    # ---- Sheet 6b：已答問題（2026-09-10 WP2b-d：老闆已一句答埋、或已補
    #      資料關閉嘅問題，逐條記低答案，唔再喺「待老闆決定」出現）----
    ws = wb.create_sheet("已答問題")
    _header_row(ws, ["問題編號", "問題", "老闆答案／結果"])
    for q in boss_facts.get("answered_questions", []):
        question = _boss_text(q["question"]) if boss else q["question"]
        answer = _boss_text(q["answer"]) if boss else q["answer"]
        ws.append([q["ref"], question, answer])
        ws.cell(row=ws.max_row, column=2).alignment = _WRAP
        ws.cell(row=ws.max_row, column=3).alignment = _WRAP
    _autosize(ws, [10, 50, 60])

    # ---- Sheet 7：已裁定口徑 ----
    ws = wb.create_sheet("已裁定口徑")
    _header_row(ws, ["裁定"])
    for line in boss_facts["decided_rules"]:
        ws.append([_boss_text(line) if boss else line])
    _autosize(ws, [110])

    path.parent.mkdir(parents=True, exist_ok=True)
    _save_stable(wb, path)
