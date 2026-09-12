"""Header-driven layout detection for the roster sheets.

P0-執行計劃.md §3.1 / §6: 三種欄位排法，唔可以寫死欄號，靠表頭定位。

All three layouts share one invariant that this module leans on: the "Name"
header cell in the header sub-row sits in the SAME COLUMN as the "Date"
label cell in the row above it. Everything else (broker code, time, rank,
staff no, tel, property name / name_en) is at a fixed column *offset* from
that shared "label column", and which set of offsets applies is decided by
inspecting the header sub-row itself (never a hardcoded column letter):

- citybus: header row contains the literal "Property Name" label.
- digital_port: the header cell immediately left of "Name" is blank, and the
  one before that is "Tel" (i.e. there's an extra unlabeled English-name
  column squeezed in between Tel and Name).
- standard: neither of the above; Tel sits directly left of Name.
"""
import datetime
from dataclasses import dataclass, field
from typing import Dict, List, Optional, Tuple

from openpyxl.worksheet.worksheet import Worksheet

STANDARD_OFFSETS = {
    "broker_code": -5,
    "time_slot": -4,
    "rank": -3,
    "staff_no": -2,
    "tel": -1,
}
DIGITAL_PORT_OFFSETS = {
    "broker_code": -6,
    "time_slot": -5,
    "rank": -4,
    "staff_no": -3,
    "tel": -2,
    "name_en": -1,
}
CITYBUS_OFFSETS = {
    "broker_code": -6,
    "time_slot": -5,
    "rank": -4,
    "sub_site": -3,
    "staff_no": -2,
    "tel": -1,
}

LAYOUT_OFFSETS = {
    "standard": STANDARD_OFFSETS,
    "digital_port": DIGITAL_PORT_OFFSETS,
    "citybus": CITYBUS_OFFSETS,
}


@dataclass
class ColumnMapping:
    date_row: int
    header_row: int
    label_col: int  # == Name column, for every layout
    layout: str
    date_start_col: int
    date_end_col: int
    date_values: Dict[int, datetime.date]  # col -> the actual date (day 1..31)
    trailing_cols: List[Tuple[int, str]]  # [(col, label), ...] right after dates

    def field_col(self, field_name: str) -> Optional[int]:
        offsets = LAYOUT_OFFSETS[self.layout]
        if field_name not in offsets:
            return None
        return self.label_col + offsets[field_name]


def _cell(ws: Worksheet, row: int, col: int):
    return ws.cell(row=row, column=col).value


def find_date_headers(ws: Worksheet) -> List[Tuple[int, int]]:
    """Find every (row, col) where a 'Date' header cell starts a block's
    column mapping. Verified by checking the very next cell is a real date.
    """
    headers = []
    max_row = ws.max_row
    max_col = ws.max_column
    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            v = _cell(ws, r, c)
            if v == "Date":
                nxt = _cell(ws, r, c + 1)
                if isinstance(nxt, datetime.datetime):
                    headers.append((r, c))
                break  # one Date header per row at most
    return headers


def detect_layout(ws: Worksheet, header_row: int, label_col: int) -> str:
    # citybus marker: 'Property Name' literally present in the header row
    for c in range(1, label_col + 1):
        if _cell(ws, header_row, c) == "Property Name":
            return "citybus"
    left1 = _cell(ws, header_row, label_col - 1)
    left2 = _cell(ws, header_row, label_col - 2) if label_col - 2 >= 1 else None
    if (left1 is None or str(left1).strip() == "") and left2 == "Tel":
        return "digital_port"
    return "standard"


def build_mapping(ws: Worksheet, date_row: int, label_col: int) -> ColumnMapping:
    date_start_col = label_col + 1
    date_values: Dict[int, datetime.date] = {}
    c = date_start_col
    while True:
        v = _cell(ws, date_row, c)
        if isinstance(v, datetime.datetime):
            date_values[c] = v.date()
            c += 1
        else:
            break
    date_end_col = c - 1 if date_values else date_start_col - 1

    header_row = date_row + 1
    layout = detect_layout(ws, header_row, label_col)

    trailing_cols: List[Tuple[int, str]] = []
    max_col = ws.max_column
    c = date_end_col + 1
    blank_run = 0
    first = True
    while c <= max_col:
        label = _cell(ws, header_row, c)
        label_str = str(label).strip() if label is not None else ""
        if first:
            # First trailing column is always the per-row monthly total,
            # even though its header cell is normally blank.
            trailing_cols.append((c, label_str or "月總更數"))
            first = False
            blank_run = 0
            c += 1
            continue
        if label_str == "":
            blank_run += 1
            if blank_run >= 2:
                break
            c += 1
            continue
        blank_run = 0
        trailing_cols.append((c, label_str))
        c += 1

    return ColumnMapping(
        date_row=date_row,
        header_row=header_row,
        label_col=label_col,
        layout=layout,
        date_start_col=date_start_col,
        date_end_col=date_end_col,
        date_values=date_values,
        trailing_cols=trailing_cols,
    )
